import csv
import json
import sys
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from metrics import analyse, DataError, cents
from review import build, extract, inspect_book, load, save, verify, digest
from make_fixture import create


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.raw=load(ROOT/'examples/demo.json')

    def test_accounting_conservation_and_known_answers(self):
        d=analyse(self.raw);t=d['totals']
        self.assertEqual(t['total_current'],47100000)
        self.assertEqual(t['total_prior'],35200000)
        self.assertEqual(t['total_delta'],11900000)
        for r in d['rows']:
            self.assertEqual(r['self_current']+r['dist_current'],r['total_current'])
        self.assertEqual(d['stats']['gross_gain']+d['stats']['gross_loss'],t['total_delta'])
        self.assertAlmostEqual(t['spend_effect']+t['ratio_effect'],t['self_delta'],places=4)
        self.assertEqual(sum(c['delta'] for c in d['channels']),t['total_delta'])

    def test_rounding_is_decimal(self):
        self.assertEqual(cents('1.005','x'),101)
        self.assertEqual(cents('-1.005','x'),-101)

    def test_missing_and_nonfinite_not_zero(self):
        for v in (None,True,float('nan'),float('inf'),'n/a'):
            with self.subTest(v=v),self.assertRaises(DataError):cents(v,'x')

    def test_duplicate_account_fails(self):
        self.raw['accounts'].append(deepcopy(self.raw['accounts'][0]))
        with self.assertRaisesRegex(DataError,'duplicate'):analyse(self.raw)

    def test_channel_and_source_checks_fail_independently(self):
        self.raw['accounts'][0]['channels'][0]['current']+=1
        with self.assertRaisesRegex(DataError,'reconcile'):analyse(self.raw)
        self.raw=load(ROOT/'examples/demo.json');self.raw['checks']['total_current']+=2
        with self.assertRaisesRegex(DataError,'summary'):analyse(self.raw)

    def test_zero_denominator_is_null(self):
        self.raw['accounts'][0]['spend']['current']=0
        d=analyse(self.raw)
        self.assertIsNone(d['rows'][0]['roi_current'])
        self.assertIsNone(d['rows'][0]['roi_delta'])

    def test_negative_distribution_preserved(self):
        self.raw.pop('checks')
        self.raw['accounts'][0]['self']['current']=200000
        d=analyse(self.raw)
        self.assertEqual(d['rows'][0]['dist_current'],-3000000)
        self.assertTrue(any('分销为负' in w for w in d['warnings']))

    def test_negative_total_hides_concentration(self):
        self.raw.pop('checks');self.raw['accounts'][0]['total']['current']=-1
        for r in self.raw['accounts']:r.pop('channels')
        d=analyse(self.raw)
        self.assertIsNone(d['stats']['cr1'])
        self.assertTrue(all(r['share'] is None for r in d['rows']))

    def test_month_end_and_mtd_confirmation(self):
        self.raw['meta']['target_context']['as_of']='2030-04-30'
        self.assertIsNone(analyse(self.raw)['pace']['daily_required'])
        self.raw['meta']['target_context']['actual_is_month_to_date']=False
        with self.assertRaisesRegex(DataError,'month-to-date'):analyse(self.raw)

    def test_mismatched_period_and_history_isolation(self):
        self.raw['meta']['baseline']['end']='2030-03-14'
        self.raw['history'][0]['rows'][0]['values']['clicks']['current']=1.5
        d=analyse(self.raw)
        self.assertTrue(any('天数不同' in w for w in d['warnings']))
        self.assertIsNone(d['history'][0]['rows'][0]['values']['clicks']['current'])
        self.assertIsNone(d['history'][0]['rows'][0]['values']['ctr']['current'])
        self.assertEqual(d['meta']['period']['start'],'2030-04-01')
        self.assertEqual(d['history'][0]['period']['start'],'2030-03-01')

    def test_head_is_dynamic_not_example_name(self):
        self.raw['accounts'][0]['name']='全新账号'
        self.assertEqual(analyse(self.raw)['stats']['head'],'全新账号')

    def test_synthetic_excel_full_pipeline(self):
        with tempfile.TemporaryDirectory() as t:
            t=Path(t);source=create(t);initial=digest(source)
            inspect_book(source,t/'inspection.json')
            extract(source,ROOT/'examples/mapping.json',t/'input.json')
            d=load(t/'input.json')
            self.assertEqual(d['accounts'][0]['name'],'示例青禾')
            self.assertEqual(len(d['accounts'][0]['channels']),3)
            self.assertTrue(d['meta']['extraction_source_unchanged'])
            build(t/'input.json',t/'out');verify(t/'out')
            self.assertEqual(digest(source),initial)
            self.assertFalse(load(t/'out/validation.json')['browser_verified'])

    def test_alias_does_not_mask_mismatched_totals(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as t:
            t=Path(t);source=create(t);wb=load_workbook(source)
            wb['渠道']['B2']=1;wb.save(source);wb.close()
            with self.assertRaisesRegex(DataError,'Join total mismatch'):extract(source,ROOT/'examples/mapping.json',t/'out.json')

    def test_uncached_formula_is_not_zero(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as t:
            t=Path(t);source=create(t);wb=load_workbook(source)
            wb['经营']['B2']='=1+1';wb.save(source);wb.close()
            with self.assertRaisesRegex(DataError,'cached value'):extract(source,ROOT/'examples/mapping.json',t/'out.json')

    def test_untrusted_text_and_artifact_tampering(self):
        self.raw['accounts'][0]['name']='=HYPERLINK("https://example.invalid")</script><img onerror=alert(1)>'
        with tempfile.TemporaryDirectory() as t:
            t=Path(t);save(t/'input.json',self.raw);build(t/'input.json',t/'out')
            html=(t/'out/report.html').read_text('utf-8')
            self.assertNotIn('</script><img onerror',html)
            self.assertIn('\\u003c/script\\u003e',html)
            with (t/'out/comparison.csv').open(encoding='utf-8-sig',newline='') as f:
                rows=list(csv.reader(f));self.assertTrue(rows[1][0].startswith("'="))
            with (t/'out/report.html').open('a',encoding='utf-8') as f:f.write('tampered')
            with self.assertRaisesRegex(DataError,'hash mismatch'):verify(t/'out')

    def test_optional_sections_can_be_absent(self):
        for k in ['history','targets','commentary']:self.raw.pop(k,None)
        self.raw['meta'].pop('target_context')
        for r in self.raw['accounts']:r.pop('channels');r.pop('target')
        d=analyse(self.raw)
        self.assertEqual(d['history'],[]);self.assertIsNone(d['pace'])
        self.assertFalse(d['validation']['channels_reconciled'])

    def test_next_period_rebuild_replaces_old_dates_accounts_and_results(self):
        with tempfile.TemporaryDirectory() as t:
            t=Path(t);save(t/'input.json',self.raw);build(t/'input.json',t/'out')
            self.raw['meta']['period']={'start':'2030-05-01','end':'2030-05-15'}
            self.raw['meta']['baseline']={'start':'2030-04-01','end':'2030-04-15'}
            self.raw['meta'].pop('target_context')
            for key in ('checks','targets','history','issues','commentary'):self.raw.pop(key,None)
            for i,r in enumerate(self.raw['accounts']):
                r['name']=f'下一期账号{i+1}';r.pop('target')
                r['channels'][0]['current']-=100;r['channels'][0]['prior']-=100
                r['channels'].append({'name':'新增渠道','current':100,'prior':100})
            extra=deepcopy(self.raw['accounts'][0]);extra['name']='新入榜头部'
            extra['total']['current']=420000
            extra['channels'][0]['current']+=250000
            self.raw['accounts'].append(extra)
            save(t/'input.json',self.raw);build(t/'input.json',t/'out')
            d=load(t/'out/analysis.json');page=(t/'out/report.html').read_text('utf-8')
            self.assertEqual(d['stats']['count'],9)
            self.assertEqual(d['stats']['head'],'新入榜头部')
            self.assertEqual(d['totals']['total_current'],89100000)
            self.assertEqual(len(d['channels']),4)
            self.assertNotIn('2030-03-01',page)
            self.assertNotIn('示例青禾',page)
            self.assertIn('2030-05-01',page)


if __name__=='__main__':unittest.main()
