#!/usr/bin/env python3
"""Portable, offline business-review CLI. No model or service dependency."""
import argparse
import base64
import csv
import hashlib
import html
from io import BytesIO
import json
from pathlib import Path
import sys

from metrics import analyse, DataError

ROOT = Path(__file__).resolve().parent.parent


def load(path):
    return json.loads(Path(path).read_text('utf-8-sig'))


def save(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def inspect_book(workbook, output):
    from openpyxl import load_workbook
    wb = load_workbook(workbook, read_only=True, data_only=False)
    cached = load_workbook(workbook, read_only=True, data_only=True)
    result = dict(source=Path(workbook).name, sha256=digest(workbook), sheets=[])
    for ws in wb:
        sheet = dict(name=ws.title, rows=ws.max_row, columns=ws.max_column, cells=[])
        # The inspection is deliberate and complete, not an arbitrary first-row sample.
        for row in ws:
            for cell in row:
                if cell.value is not None:
                    value = cached[ws.title][cell.coordinate].value if cell.data_type == 'f' else cell.value
                    sheet['cells'].append(dict(cell=cell.coordinate, value=value, formula=cell.value if cell.data_type == 'f' else None))
        result['sheets'].append(sheet)
    wb.close(); cached.close()
    # Excel dates may be datetime objects in inspection; do not silently coerce monetary data.
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    Path(output).write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding='utf-8')


def extract(workbook, mapping, output):
    from openpyxl import load_workbook
    m = load(mapping)
    initial = digest(workbook)
    # Read a snapshot so validation failures cannot leave the user's file locked.
    wb = load_workbook(BytesIO(Path(workbook).read_bytes()), read_only=True, data_only=True)
    refs = []

    def resolve(value, row=None):
        if isinstance(value, dict) and '$cell' in value:
            sheet, address = value['$cell'].split('!', 1)
            address = address.replace('{row}', str(row))
            cell = wb[sheet][address]
            v = cell.value
            refs.append(f'{sheet}!{address}')
            if v is None and not value.get('nullable', False):
                raise DataError(f'{sheet}!{address}: blank or formula has no cached value; recalculate source in Excel')
            if v is not None and 'scale' in value:
                if isinstance(v,bool) or not isinstance(v,(int,float)):
                    raise DataError(f'{sheet}!{address}: numeric scale applied to non-number')
                v *= value['scale']
            return v
        if isinstance(value, dict):
            return {k:resolve(v,row) for k,v in value.items()}
        if isinstance(value,list):
            return [resolve(v,row) for v in value]
        return value

    raw = resolve(m.get('base',{}))
    raw.setdefault('schema_version',1)
    raw.setdefault('accounts',[])
    for block in m.get('account_blocks',[]):
        for row in block['rows']:
            refs.clear()
            account = resolve(block['record'],row)
            account['source_refs'] = list(dict.fromkeys(refs))
            raw['accounts'].append(account)
    by_name = {a['name']:a for a in raw['accounts']}
    for join in m.get('joins',[]):
        seen = set()
        for row in join['rows']:
            refs.clear()
            rec = resolve(join['record'],row)
            name = join.get('aliases',{}).get(rec['name'],rec['name'])
            if name not in by_name or name in seen:
                raise DataError(f'Join unmatched or duplicate account: {name}')
            seen.add(name)
            a = by_name[name]
            for p, value in rec.pop('check_total',{}).items():
                from metrics import cents
                if abs(cents(a['total'][p],name)-cents(value,name)) > 1:
                    raise DataError(f'Join total mismatch: {name} {p}')
            for key,value in rec.items():
                if key == 'name':
                    continue
                if key in a:
                    raise DataError(f'Join cannot overwrite existing field: {name}.{key}')
                a[key] = value
            a['source_refs'].extend(refs)
        if join.get('require_all',True) and seen != set(by_name):
            raise DataError('Join omitted accounts; map missing rows explicitly')
    wb.close()
    if initial != digest(workbook):
        raise DataError('Source workbook changed during extraction')
    raw.setdefault('meta',{})['source'] = Path(workbook).name
    raw['meta']['source_sha256'] = initial
    raw['meta']['extraction_source_unchanged'] = True
    analyse(raw)
    save(output,raw)


CSV_FIELDS = [('name','账号'),('agency','代理商'),('total_current','本期总 GMV'),('total_prior','上期总 GMV'),
              ('self_current','本期自营'),('self_prior','上期自营'),('dist_current','本期分销'),('dist_prior','上期分销'),
              ('spend_current','本期消耗'),('spend_prior','上期消耗'),('total_delta','总 GMV 增减'),
              ('total_growth','总 GMV 增长率'),('roi_current','本期自营 GMV/消耗'),('roi_prior','上期自营 GMV/消耗')]


def safe_cell(value):
    # Text imported from a workbook must not become an executable spreadsheet formula.
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')):
        return "'"+value
    return value


def output_csv(d, path):
    with Path(path).open('w',encoding='utf-8-sig',newline='') as f:
        w = csv.writer(f)
        w.writerow([label+'（元）' if key not in ('name','agency','total_growth','roi_current','roi_prior') else label for key,label in CSV_FIELDS])
        for r in d['rows']:
            w.writerow([safe_cell(r[k]) if k in ('name','agency') else r[k] if k in ('total_growth','roi_current','roi_prior') else r[k]/100 for k,_ in CSV_FIELDS])


def output_xlsx(d,path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.workbook.properties import CalcProperties
    wb = Workbook(); ws = wb.active; ws.title='分销和自营数据对比'
    ws.append(['账号','本期总 GMV','上期总 GMV','本期自营','上期自营','本期分销','上期分销','本期消耗','上期消耗','总 GMV 增减','总 GMV 增长率','本期自营/消耗','上期自营/消耗','代理商'])
    for r in d['rows']:
        i = ws.max_row+1
        ws.append([safe_cell(r['name']),r['total_current']/100,r['total_prior']/100,r['self_current']/100,r['self_prior']/100,
                   f'=B{i}-D{i}',f'=C{i}-E{i}',r['spend_current']/100,r['spend_prior']/100,f'=B{i}-C{i}',
                   f'=IF(C{i}>0,J{i}/C{i},"")',f'=IF(H{i}>0,D{i}/H{i},"")',f'=IF(I{i}>0,E{i}/I{i},"")',safe_cell(r['agency'])])
    last=ws.max_row; total=last+1
    ws.append(['合计']+[f'=SUM({c}2:{c}{last})' for c in 'BCDEFGHIJ']+[f'=IF(C{total}>0,J{total}/C{total},"")',f'=IF(H{total}>0,D{total}/H{total},"")',f'=IF(I{total}>0,E{total}/I{total},"")',''])
    for row in ws:
        for c in row:
            c.alignment=Alignment(vertical='center')
            if 2<=c.column<=10:c.number_format='#,##0.00;[Red]-#,##0.00'
            if c.column==11:c.number_format='0.0%'
            if c.column in (12,13):c.number_format='0.00'
            if c.row in (1,total):
                c.fill=PatternFill('solid',fgColor='1E3282');c.font=Font(color='FFFFFF',bold=True)
    from openpyxl.utils import get_column_letter
    for i in range(1,15):ws.column_dimensions[get_column_letter(i)].width=25 if i in (1,14) else 19
    ws.freeze_panes='B2';ws.auto_filter.ref=f'A1:N{last}'
    notes=wb.create_sheet('口径与核验')
    for pair in [('本期',d['meta']['period']['label']),('上期',d['meta']['baseline']['label']),('金额单位','元'),('来源',d['meta'].get('source','规范化输入')),('公式缓存','公式由 Excel / WPS 打开时重算；预计算值见 CSV 与 HTML。'),('效率说明','自营 GMV / 消耗，不等于利润或广告归因回报。')]:notes.append(pair)
    for w in d['warnings']:notes.append(['核验提示',str(w)])
    notes.column_dimensions['A'].width=20;notes.column_dimensions['B'].width=100
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True)
    wb.save(path)


def build(input_path,out):
    before=digest(input_path)
    d=analyse(load(input_path))
    output=Path(out);output.mkdir(parents=True,exist_ok=True)
    # Escape script terminators, not just HTML text; account labels are untrusted data.
    payload=json.dumps(d,ensure_ascii=False,separators=(',',':'),allow_nan=False).replace('<','\\u003c').replace('>','\\u003e').replace('&','\\u0026')
    template=(ROOT/'assets/report.html').read_text('utf-8')
    replacements={'TITLE':html.escape(d['meta']['title']), 'CSS':(ROOT/'assets/report.css').read_text('utf-8'),
                  'LOGO':base64.b64encode((ROOT/'assets/brand/company-logo.png').read_bytes()).decode(),
                  'ECHARTS':(ROOT/'assets/vendor/echarts.min.js').read_text('utf-8'), 'DATA':payload,
                  'APP':(ROOT/'assets/report.js').read_text('utf-8')}
    for k,v in replacements.items():template=template.replace('{{'+k+'}}',v)
    (output/'report.html').write_text(template,encoding='utf-8')
    output_csv(d,output/'comparison.csv');output_xlsx(d,output/'comparison.xlsx')
    save(output/'analysis.json',d)
    if digest(input_path)!=before:raise DataError('Input changed during build')
    report=dict(status='passed',input_sha256=before,input_unchanged=True,checks=d['validation'],warnings=d['warnings'],
                browser_verified=False,artifacts={p:digest(output/p) for p in ('report.html','comparison.csv','comparison.xlsx','analysis.json')})
    save(output/'validation.json',report)
    verify(output)


def verify(out):
    from openpyxl import load_workbook
    out=Path(out);v=load(out/'validation.json');d=load(out/'analysis.json')
    for name,sha in v['artifacts'].items():
        if digest(out/name)!=sha:raise DataError(f'Artifact hash mismatch: {name}')
    report=(out/'report.html').read_text('utf-8')
    if any('{{'+k+'}}' in report for k in ('TITLE','DATA','LOGO','APP','CSS','ECHARTS')):raise DataError('Unexpanded template')
    if '<script src=' in report:raise DataError('External script dependency')
    with (out/'comparison.csv').open(encoding='utf-8-sig',newline='') as f:
        rows=list(csv.reader(f))
    if len(rows)!=len(d['rows'])+1:raise DataError('CSV row count mismatch')
    wb=load_workbook(out/'comparison.xlsx',read_only=True,data_only=False)
    ws=wb.worksheets[0]
    if ws.max_row!=len(d['rows'])+2:raise DataError('Workbook row count mismatch')
    for i,r in enumerate(d['rows'],2):
        if ws[f'F{i}'].value!=f'=B{i}-D{i}' or abs(ws[f'B{i}'].value*100-r['total_current'])>0.01:
            raise DataError('Workbook formula/source mismatch')
    wb.close()
    print(f'PASS: {len(d["rows"])} accounts; offline HTML, CSV, formula workbook and hashes verified. Browser QA is separate.')


def main():
    p=argparse.ArgumentParser(description=__doc__);sub=p.add_subparsers(dest='command',required=True)
    x=sub.add_parser('inspect');x.add_argument('--workbook',required=True);x.add_argument('--out',required=True)
    x=sub.add_parser('extract');x.add_argument('--workbook',required=True);x.add_argument('--mapping',required=True);x.add_argument('--out',required=True)
    x=sub.add_parser('build');x.add_argument('--input',required=True);x.add_argument('--out',required=True)
    x=sub.add_parser('verify');x.add_argument('--out',required=True)
    a=p.parse_args()
    try:
        if a.command=='inspect':inspect_book(a.workbook,a.out)
        elif a.command=='extract':extract(a.workbook,a.mapping,a.out)
        elif a.command=='build':build(a.input,a.out)
        else:verify(a.out)
    except (DataError,KeyError,ValueError) as e:
        print(f'ERROR: {e}',file=sys.stderr);return 2
    return 0


if __name__=='__main__':
    raise SystemExit(main())
